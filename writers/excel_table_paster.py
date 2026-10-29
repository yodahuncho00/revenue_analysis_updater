import os
import re
import sys
import time
from decimal import Decimal, InvalidOperation
from typing import Dict, Tuple

import pywintypes
import win32clipboard
import win32con
import win32com.client
import openpyxl
from openpyxl.utils import get_column_letter

from extractors.base import BalanceTableCapture


class ExcelTablePaster:
    """Paste captured HTML tables into the working workbook using Excel's HTML clipboard."""

    COM_RETRY_ATTEMPTS = 40
    COM_RETRY_DELAY_SECONDS = 0.25
    RPC_E_CALL_REJECTED = -2147418111
    COM_DISCONNECTED_ERRORS = {-2147417848, -2147023174}

    AMOUNT_COLUMNS = {
        "BO_merchant_balance": (3, 6),
        "BO_provider_balance": (5, 8),
    }
    HELPER_FORMAT_COLUMNS = {
        "BO_merchant_balance": (9, 21),
        "BO_provider_balance": (11, 20),
    }
    HELPER_NUMBER_FORMAT_LOCAL = '_(* # ##0_);_(* (# ##0);_(* "-"??_);_(@_)'
    FIRST_FORMULA_ROW = 6
    XL_LEFT = -4131
    XL_RIGHT = -4152
    XL_NONE = -4142

    @staticmethod
    def _clipboard_html(table_html: str) -> bytes:
        start_marker = "<!--StartFragment-->"
        end_marker = "<!--EndFragment-->"
        html = f"<html><body>{start_marker}{table_html}{end_marker}</body></html>"
        html_bytes = html.encode("utf-8")
        header_template = (
            "Version:0.9\r\n"
            "StartHTML:{:010d}\r\n"
            "EndHTML:{:010d}\r\n"
            "StartFragment:{:010d}\r\n"
            "EndFragment:{:010d}\r\n"
        )
        header = header_template.format(0, 0, 0, 0).encode("ascii")
        start_html = len(header)
        start_fragment = start_html + html_bytes.index(start_marker.encode("ascii")) + len(
            start_marker
        )
        end_fragment = start_html + html_bytes.index(end_marker.encode("ascii"))
        end_html = start_html + len(html_bytes)
        header = header_template.format(
            start_html, end_html, start_fragment, end_fragment
        ).encode("ascii")
        return header + html_bytes

    @staticmethod
    def _cell_matrix(
        capture: BalanceTableCapture, *, normalize_line_breaks: bool = False
    ):
        matrix = []
        active_rowspans: Dict[int, int] = {}
        for row in capture.rows:
            cells = row.get("cells", [])
            spans = row.get("spans", [])
            expanded_cells = [""] * capture.column_count
            occupied_columns = set(active_rowspans)
            next_rowspans = {
                column: remaining - 1
                for column, remaining in active_rowspans.items()
                if remaining > 1
            }
            column = 0

            for index, cell in enumerate(cells):
                while column in occupied_columns:
                    column += 1
                span = spans[index] if index < len(spans) else {}
                column_span = max(1, span.get("col_span", 1))
                row_span = max(1, span.get("row_span", 1))
                if column < len(expanded_cells):
                    expanded_cells[column] = (
                        " ".join(cell.split())
                        if normalize_line_breaks
                        else cell
                    )

                for occupied_column in range(
                    column, min(column + column_span, len(expanded_cells))
                ):
                    occupied_columns.add(occupied_column)
                    if row_span > 1:
                        next_rowspans[occupied_column] = max(
                            next_rowspans.get(occupied_column, 0), row_span - 1
                        )
                column += column_span

            matrix.append(tuple(expanded_cells))
            active_rowspans = next_rowspans

        return tuple(matrix)

    @classmethod
    def _plain_text(cls, capture: BalanceTableCapture) -> str:
        return "\r\n".join(
            "\t".join(row)
            for row in cls._cell_matrix(capture, normalize_line_breaks=True)
        )

    @classmethod
    def _set_clipboard(cls, capture: BalanceTableCapture) -> None:
        html_format = win32clipboard.RegisterClipboardFormat("HTML Format")
        last_error = None
        for _ in range(20):
            try:
                win32clipboard.OpenClipboard()
                break
            except pywintypes.error as exc:
                last_error = exc
                time.sleep(0.05)
        else:
            raise RuntimeError("Could not open the Windows clipboard for Excel table paste.") from last_error

        try:
            win32clipboard.EmptyClipboard()
            win32clipboard.SetClipboardData(
                html_format,
                cls._clipboard_html(
                    re.sub(r"<br\b[^>]*>", " ", capture.html, flags=re.IGNORECASE)
                ),
            )
            win32clipboard.SetClipboardData(
                win32con.CF_UNICODETEXT, cls._plain_text(capture)
            )
        finally:
            win32clipboard.CloseClipboard()

    @classmethod
    def _retry_excel_call(cls, operation):
        for attempt in range(cls.COM_RETRY_ATTEMPTS):
            try:
                return operation()
            except pywintypes.com_error as exc:
                if exc.args[0] != cls.RPC_E_CALL_REJECTED:
                    raise
                if attempt == cls.COM_RETRY_ATTEMPTS - 1:
                    raise
                time.sleep(cls.COM_RETRY_DELAY_SECONDS)

    @classmethod
    def _paste_html_table(cls, worksheet, top_left, capture) -> None:
        for attempt in range(3):
            cls._set_clipboard(capture)
            cls._retry_excel_call(lambda: worksheet.Select())
            cls._retry_excel_call(lambda: top_left.Select())
            try:
                cls._retry_excel_call(
                    lambda: worksheet.Paste(Destination=top_left)
                )
                return
            except pywintypes.com_error as exc:
                details = exc.args[2] if len(exc.args) > 2 else ()
                is_excel_paste_failure = (
                    isinstance(details, tuple)
                    and len(details) > 5
                    and details[1] == "Microsoft Excel"
                    and details[-1] == -2146827284
                )
                if not is_excel_paste_failure or attempt == 2:
                    raise
                time.sleep(cls.COM_RETRY_DELAY_SECONDS)

    @staticmethod
    def _parse_amount(value: str):
        text = value.strip().replace("\u00a0", "").replace("\u202f", "").replace(" ", "")
        if not text:
            return None
        if text in {"—", "–", "-"}:
            return text

        if "," in text and "." in text:
            text = text.replace(",", "")
        elif "," in text:
            if re.fullmatch(r"[+-]?\d{1,3}(?:,\d{3})+", text):
                text = text.replace(",", "")
            else:
                text = text.replace(",", ".")

        try:
            return float(Decimal(text))
        except InvalidOperation as exc:
            raise ValueError(
                f"Unexpected balance amount from Backoffice: {value!r}."
            ) from exc

    @classmethod
    def _write_numeric_amounts(
        cls, worksheet, sheet_name, capture, top_row, first_column
    ) -> None:
        amount_columns = cls.AMOUNT_COLUMNS.get(sheet_name)

        if amount_columns is not None:
            start_index, end_index = amount_columns
            values = []
            for row in capture.rows:
                cells = row.get("cells", [])
                if "headers-row" in row.get("row_class", "").split():
                    values.append(
                        tuple(
                            cells[column] if column < len(cells) else ""
                            for column in range(start_index, end_index)
                        )
                    )
                    continue
                values.append(
                    tuple(
                        cls._parse_amount(str(cells[column]))
                        if column < len(cells) and str(cells[column]).strip()
                        else ""
                        for column in range(start_index, end_index)
                    )
                )

            first_amount_column = get_column_letter(first_column + start_index)
            last_amount_column = get_column_letter(first_column + end_index - 1)
            amount_range = cls._retry_excel_call(
                lambda: worksheet.Range(
                    f"{first_amount_column}{top_row}:"
                    f"{last_amount_column}{top_row + capture.row_count - 1}"
                )
            )
            cls._retry_excel_call(
                lambda: setattr(amount_range, "Value2", tuple(values))
            )
            cls._retry_excel_call(
                lambda: setattr(
                    amount_range,
                    "NumberFormatLocal",
                    "# ##0,00;-# ##0,00;0",
                )
            )

    @classmethod
    def _apply_captured_row_heights(
        cls, worksheet, top_row: int, capture: BalanceTableCapture
    ) -> None:
        runs = []
        for offset, row in enumerate(capture.rows):
            height = row.get("height")
            if not isinstance(height, (int, float)) or height <= 0:
                continue
            row_number = top_row + offset
            height_points = round(height * 0.75, 2)
            if (
                runs
                and runs[-1][1] + 1 == row_number
                and runs[-1][2] == height_points
            ):
                runs[-1][1] = row_number
            else:
                runs.append([row_number, row_number, height_points])

        for first_row, last_row, height_points in runs:
            cls._retry_excel_call(
                lambda: setattr(
                    worksheet.Rows(f"{first_row}:{last_row}"),
                    "RowHeight",
                    height_points,
                )
            )

    @classmethod
    def _write_captured_cells(
        cls, worksheet, top_row: int, first_column: int,
        capture: BalanceTableCapture
    ) -> None:
        bottom_right = cls._retry_excel_call(
            lambda: worksheet.Cells(
                top_row + capture.row_count - 1,
                first_column + capture.column_count - 1,
            )
        )
        target_range = cls._retry_excel_call(
            lambda: worksheet.Range(
                worksheet.Cells(top_row, first_column), bottom_right
            )
        )
        values = cls._cell_matrix(capture)
        cls._retry_excel_call(
            lambda: setattr(target_range, "Value2", values)
        )

        multiline_cells = [
            (top_row + row_offset, first_column + column_offset)
            for row_offset, row in enumerate(values)
            for column_offset, value in enumerate(row)
            if "\n" in value or "\r" in value
        ]
        multiline_by_column: Dict[int, list[int]] = {}
        for row, column in multiline_cells:
            multiline_by_column.setdefault(column, []).append(row)
        for column, rows in multiline_by_column.items():
            row_runs = []
            for row in rows:
                if row_runs and row == row_runs[-1][1] + 1:
                    row_runs[-1][1] = row
                else:
                    row_runs.append([row, row])
            for first_row, last_row in row_runs:
                cell_range = cls._retry_excel_call(
                    lambda: worksheet.Range(
                        worksheet.Cells(first_row, column),
                        worksheet.Cells(last_row, column),
                    )
                )
                cls._retry_excel_call(
                    lambda: setattr(cell_range, "WrapText", True)
                )

    @classmethod
    def paste_tables(
        cls,
        workbook_path: str,
        tables: Dict[str, Tuple[BalanceTableCapture, str]],
    ) -> None:
        # A disconnected COM proxy cannot be revived by retrying the same
        # worksheet/workbook object. Start a fresh Excel instance and replay
        # the deterministic table replacement instead.
        for attempt in range(3):
            try:
                return cls._paste_tables_once(workbook_path, tables)
            except pywintypes.com_error as exc:
                error_code = exc.args[0] if exc.args else None
                if (
                    error_code not in cls.COM_DISCONNECTED_ERRORS
                    or attempt == 2
                ):
                    raise
                time.sleep(cls.COM_RETRY_DELAY_SECONDS * (attempt + 1))

    @classmethod
    def _paste_tables_once(
        cls,
        workbook_path: str,
        tables: Dict[str, Tuple[BalanceTableCapture, str]],
    ) -> None:
        if not tables:
            return

        source_workbook = openpyxl.load_workbook(
            workbook_path, read_only=False, data_only=False
        )
        merged_rows = {}
        start_positions = {}
        try:
            for sheet_name, (_, start_cell) in tables.items():
                worksheet = source_workbook[sheet_name]
                start = worksheet[start_cell]
                start_positions[sheet_name] = (start.row, start.column)
                last_column = start.column + tables[sheet_name][0].column_count - 1
                merged_rows[sheet_name] = max(
                    (
                        cell_range.max_row
                        for cell_range in worksheet.merged_cells.ranges
                        if cell_range.min_col <= last_column
                        and cell_range.max_col >= start.column
                    ),
                    default=0,
                )
        finally:
            source_workbook.close()

        excel = win32com.client.gencache.EnsureDispatch(
            win32com.client.DispatchEx("Excel.Application")
        )
        cls._retry_excel_call(lambda: setattr(excel, "Visible", False))
        cls._retry_excel_call(lambda: setattr(excel, "DisplayAlerts", False))
        workbook = None
        saved = False
        try:
            workbook = cls._retry_excel_call(
                lambda: excel.Workbooks.Open(
                    os.path.abspath(workbook_path),
                    UpdateLinks=0,
                    ReadOnly=False,
                )
            )
            for sheet_name, (capture, start_cell) in tables.items():
                worksheet = cls._retry_excel_call(
                    lambda: workbook.Worksheets(sheet_name)
                )
                top_left = cls._retry_excel_call(
                    lambda: worksheet.Range(start_cell)
                )
                top_row, left_column = start_positions[sheet_name]
                if sheet_name == "BO_provider_balance" and top_row == 6:
                    empty_header = cls._retry_excel_call(
                        lambda: worksheet.Range("A5:H5")
                    )
                    cls._retry_excel_call(
                        lambda: empty_header.ClearContents()
                    )
                    helper_header = cls._retry_excel_call(
                        lambda: worksheet.Range("K5:T5")
                    )
                    cls._retry_excel_call(
                        lambda: helper_header.ClearContents()
                    )
                last_existing_row = max(
                    top_row + capture.row_count - 1,
                    merged_rows[sheet_name],
                )
                for column in range(
                    left_column,
                    left_column + capture.column_count,
                ):
                    existing_last = cls._retry_excel_call(
                        lambda: worksheet.Cells(
                            worksheet.Rows.Count, column
                        ).End(-4162).Row
                    )
                    last_existing_row = max(last_existing_row, existing_last)

                clear_range = cls._retry_excel_call(
                    lambda: worksheet.Range(
                        top_left,
                        worksheet.Cells(
                            last_existing_row,
                            left_column + capture.column_count - 1,
                        ),
                    ),
                )
                cls._retry_excel_call(lambda: clear_range.UnMerge())
                cls._retry_excel_call(lambda: clear_range.ClearContents())
                cls._paste_html_table(worksheet, top_left, capture)
                worksheet = cls._retry_excel_call(
                    lambda: workbook.Worksheets(sheet_name)
                )
                pasted_last_row = top_row + capture.row_count - 1
                bottom_right = cls._retry_excel_call(
                    lambda: worksheet.Cells(
                        pasted_last_row,
                        left_column + capture.column_count - 1,
                    )
                )
                pasted_range = cls._retry_excel_call(
                    lambda: worksheet.Range(
                        worksheet.Cells(top_row, left_column), bottom_right
                    )
                )
                cls._retry_excel_call(lambda: pasted_range.UnMerge())
                cls._write_captured_cells(
                    worksheet, top_row, left_column, capture
                )
                if sheet_name == "BO_provider_balance":
                    cls._apply_captured_row_heights(
                        worksheet, top_row, capture
                    )
                if pasted_last_row < last_existing_row:
                    stale_tail = cls._retry_excel_call(
                        lambda: worksheet.Range(
                            worksheet.Cells(
                                pasted_last_row + 1, left_column
                            ),
                            worksheet.Cells(
                                last_existing_row,
                                left_column + capture.column_count - 1,
                            ),
                        ),
                    )
                    cls._retry_excel_call(lambda: stale_tail.UnMerge())
                    cls._retry_excel_call(lambda: stale_tail.ClearContents())
                    cls._retry_excel_call(
                        lambda: setattr(
                            stale_tail.Borders,
                            "LineStyle",
                            cls.XL_NONE,
                        )
                    )
                worksheet = cls._retry_excel_call(
                    lambda: workbook.Worksheets(sheet_name)
                )
                cls._write_numeric_amounts(
                    worksheet, sheet_name, capture, top_row, left_column
                )
                helper_columns = cls.HELPER_FORMAT_COLUMNS.get(sheet_name)
                if helper_columns and pasted_last_row >= cls.FIRST_FORMULA_ROW:
                    helper_range = cls._retry_excel_call(
                        lambda: worksheet.Range(
                            worksheet.Cells(
                                5, helper_columns[0]
                            ),
                            worksheet.Cells(
                                pasted_last_row, helper_columns[1]
                            ),
                        )
                    )
                    cls._retry_excel_call(
                        lambda: setattr(
                            helper_range,
                            "NumberFormatLocal",
                            cls.HELPER_NUMBER_FORMAT_LOCAL,
                        )
                    )
                if sheet_name == "BO_merchant_balance":
                    cls._retry_excel_call(
                        lambda: setattr(
                            worksheet.Rows(f"{top_row}:{last_existing_row}"),
                            "RowHeight",
                            13,
                        )
                    )
            cls._retry_excel_call(lambda: workbook.Save())
            saved = True
        finally:
            active_exception = sys.exc_info()[0] is not None
            cleanup_error = None
            try:
                if workbook is not None:
                    try:
                        cls._retry_excel_call(
                            lambda: workbook.Close(SaveChanges=saved)
                        )
                    except pywintypes.com_error as exc:
                        if exc.args[0] not in cls.COM_DISCONNECTED_ERRORS:
                            raise
                        cleanup_error = exc
            finally:
                try:
                    cls._retry_excel_call(lambda: excel.Quit())
                except pywintypes.com_error as exc:
                    if exc.args[0] not in cls.COM_DISCONNECTED_ERRORS:
                        raise
                    if cleanup_error is None:
                        cleanup_error = exc
            if cleanup_error is not None and not active_exception:
                raise cleanup_error

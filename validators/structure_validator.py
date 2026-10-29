import os
import openpyxl
from typing import Dict, Any, List

class StructureValidationError(Exception):
    pass

class WorkbookStructureValidator:
    """Pre-flight validator checking existence, accessibility, expected sheets, and header schemas."""
    
    EXPECTED_SHEETS = [
        "Control",
        "BO_extract",
        "Data",
        "BO_merchant_balance",
        "BO_provider_balance",
        "BO_settlements",
        "BO_topups"
    ]

    EXPECTED_HEADERS = {
        "BO_extract": {
            "row": 1,
            "columns": {
                1: "Period (month-end / FY = 31-Dec)",
                2: "Type (M / FY)",
                3: "Country",
                4: "Currency",
                5: "Merchant",
                6: "Payment Method",
                7: "Provider",
                8: "Deposits",
                9: "Deposits Fee",
                10: "Fee%",
                11: "Deposits Cost",
                12: "Cost%",
                13: "Withdrawals",
                14: "Withdrawals Fee",
                15: "Fee%",
                16: "Withdrawals Cost",
                17: "Cost%",
                18: "Settlement",
                19: "Topup",
                20: "Merchant Delta",
                21: "Provider Delta",
                22: "Partner Revenue"
            }
        },
        "Data": {
            "row": 4,
            "columns": {
                1: "Period (month-end / FY = 31-Dec)",
                2: "Type (M / FY)",
                3: "Country",
                4: "Currency",
                5: "Merchant",
                6: "Payment Method",
                7: "Provider account",
                8: "Deposits",
                9: "Deposits Fee",
                10: "Fee%",
                11: "Deposits Cost",
                12: "Cost%",
                13: "Withdrawals",
                14: "Withdrawals Fee",
                15: "Fee%",
                16: "Withdrawals Cost",
                17: "Cost%",
                18: "Settlement",
                19: "Topup",
                20: "Merchant Delta",
                21: "Provider Delta",
                22: "Partner Revenue"
            }
        },
        "BO_settlements": {
            "row": 1,
            "columns": {
                1: "Id",
                2: "Created",
                3: "Method",
                4: "Provider Id",
                5: "Status",
                6: "Amount",
                7: "Currency",
                8: "Country",
                9: "Merchant",
                10: "Merchant Order Id",
                11: "Note",
                12: "Year",
                13: "Month"
            }
        },
        "BO_topups": {
            "row": 1,
            "columns": {
                1: "Id",
                2: "Created",
                3: "Method",
                4: "Provider Id",
                5: "Status",
                6: "Amount",
                7: "Currency",
                8: "Country",
                9: "Merchant",
                10: "Merchant Order Id",
                11: "Note",
                12: "Year",
                13: "Month"
            }
        }
    }

    @classmethod
    def validate(cls, workbook_path: str) -> None:
        if not os.path.exists(workbook_path):
            raise StructureValidationError(f"Target workbook does not exist: {workbook_path}")

        try:
            # Check read access
            with open(workbook_path, "rb") as f:
                f.read(100)
        except PermissionError as e:
            raise StructureValidationError(f"Workbook is locked or currently open in Excel: {e}")
        except Exception as e:
            raise StructureValidationError(f"Cannot access workbook: {e}")

        # Load workbook (read-only mode is fast and doesn't mutate)
        try:
            wb = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
        except Exception as e:
            raise StructureValidationError(f"Failed to parse workbook XML: {e}")

        # Check sheets
        missing_sheets = [s for s in cls.EXPECTED_SHEETS if s not in wb.sheetnames]
        if missing_sheets:
            wb.close()
            raise StructureValidationError(f"Missing required sheets in workbook: {missing_sheets}")

        # Check headers
        for sheet_name, cfg in cls.EXPECTED_HEADERS.items():
            ws = wb[sheet_name]
            header_row = cfg["row"]
            # Read header row
            for col_idx, expected_header in cfg["columns"].items():
                val = ws.cell(header_row, col_idx).value
                val_clean = str(val).strip() if val is not None else ""
                expected_clean = expected_header.strip()
                if not val_clean.startswith(expected_clean):
                    wb.close()
                    raise StructureValidationError(
                        f"Header mismatch in {sheet_name} at row {header_row}, col {col_idx}: "
                        f"Expected starting with '{expected_clean}', got '{val_clean}'"
                    )

        # Check Control sheet label
        ws_ctrl = wb["Control"]
        as_of_found = False
        for r in range(1, 15):
            val = ws_ctrl.cell(r, 1).value
            if val and "Data as-of date (last day in the export)" in str(val):
                as_of_found = True
                break
        if not as_of_found:
            wb.close()
            raise StructureValidationError("Control sheet missing 'Data as-of date (last day in the export)' label in Column A")

        wb.close()

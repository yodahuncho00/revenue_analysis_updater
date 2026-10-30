import os
import sys
import json
import openpyxl

sys.stdout.reconfigure(encoding='utf-8')

WORKBOOK_PATH = "05. Revenue_analysis_copy.xlsx"
FIXTURES_DIR = "data/fixtures"
os.makedirs(FIXTURES_DIR, exist_ok=True)

wb = openpyxl.load_workbook(WORKBOOK_PATH, data_only=True)

# 1. Dashboard / BO_extract October rows
ws_bo = wb["BO_extract"]
dashboard_rows = []
for r in range(1088, 1230):
    dashboard_rows.append({
        "country": ws_bo.cell(r, 3).value,
        "currency": ws_bo.cell(r, 4).value,
        "merchant": ws_bo.cell(r, 5).value,
        "payment_method": ws_bo.cell(r, 6).value,
        "provider": ws_bo.cell(r, 7).value,
        "deposits": ws_bo.cell(r, 8).value,
        "deposits_fee": ws_bo.cell(r, 9).value,
        "fee_pct_dep": ws_bo.cell(r, 10).value,
        "deposits_cost": ws_bo.cell(r, 11).value,
        "cost_pct_dep": ws_bo.cell(r, 12).value,
        "withdrawals": ws_bo.cell(r, 13).value,
        "withdrawals_fee": ws_bo.cell(r, 14).value,
        "fee_pct_wd": ws_bo.cell(r, 15).value,
        "withdrawals_cost": ws_bo.cell(r, 16).value,
        "cost_pct_wd": ws_bo.cell(r, 17).value,
        "settlement": ws_bo.cell(r, 18).value,
        "topup": ws_bo.cell(r, 19).value,
        "merchant_delta": ws_bo.cell(r, 20).value,
        "provider_delta": ws_bo.cell(r, 21).value,
        "partner_revenue": ws_bo.cell(r, 22).value,
    })

with open(os.path.join(FIXTURES_DIR, "dashboard.json"), "w", encoding="utf-8") as f:
    json.dump(dashboard_rows, f, indent=2, ensure_ascii=False)
print(f"Saved {len(dashboard_rows)} dashboard fixture rows")

# 2. Merchant balance rows
ws_mb = wb["BO_merchant_balance"]
mb_rows = []
for r in range(5, ws_mb.max_row + 1):
    mb_rows.append({
        "col_a": ws_mb.cell(r, 1).value,
        "col_b": ws_mb.cell(r, 2).value,
        "col_c": ws_mb.cell(r, 3).value,
        "total": ws_mb.cell(r, 4).value,
        "pending": ws_mb.cell(r, 5).value,
        "available": ws_mb.cell(r, 6).value,
    })

with open(os.path.join(FIXTURES_DIR, "merchant_balance.json"), "w", encoding="utf-8") as f:
    json.dump(mb_rows, f, indent=2, ensure_ascii=False)
print(f"Saved {len(mb_rows)} merchant balance fixture rows")

# 3. Provider balance rows
ws_pb = wb["BO_provider_balance"]
pb_rows = []
for r in range(6, ws_pb.max_row + 1):
    pb_rows.append({
        "col_a": ws_pb.cell(r, 1).value,
        "col_b": ws_pb.cell(r, 2).value,
        "col_c": ws_pb.cell(r, 3).value,
        "on_date": ws_pb.cell(r, 5).value,
        "provider_balance": ws_pb.cell(r, 6).value,
        "calculated_balance": ws_pb.cell(r, 7).value,
        "merchant_balance": ws_pb.cell(r, 8).value,
    })

with open(os.path.join(FIXTURES_DIR, "provider_balance.json"), "w", encoding="utf-8") as f:
    json.dump(pb_rows, f, indent=2, ensure_ascii=False)
print(f"Saved {len(pb_rows)} provider balance fixture rows")

# 4. Settlements
ws_st = wb["BO_settlements"]
settle_rows = []
for r in range(max(2, ws_st.max_row - 20), ws_st.max_row + 1):
    if ws_st.cell(r, 1).value is not None:
        created_val = ws_st.cell(r, 2).value
        created_str = created_val.strftime("%Y-%m-%d %H:%M:%S") if hasattr(created_val, "strftime") else str(created_val)
        settle_rows.append({
            "id": ws_st.cell(r, 1).value,
            "created": created_str,
            "method": ws_st.cell(r, 3).value,
            "provider_id": ws_st.cell(r, 4).value,
            "status": ws_st.cell(r, 5).value,
            "amount": ws_st.cell(r, 6).value,
            "currency": ws_st.cell(r, 7).value,
            "country": ws_st.cell(r, 8).value,
            "merchant": ws_st.cell(r, 9).value,
            "merchant_order_id": ws_st.cell(r, 10).value,
            "note": ws_st.cell(r, 11).value
        })

with open(os.path.join(FIXTURES_DIR, "settlements.json"), "w", encoding="utf-8") as f:
    json.dump(settle_rows, f, indent=2, ensure_ascii=False)
print(f"Saved {len(settle_rows)} settlements fixture rows")

# 5. Topups
ws_tp = wb["BO_topups"]
topup_rows = []
for r in range(max(2, 66 - 20), 67):
    if ws_tp.cell(r, 1).value is not None:
        created_val = ws_tp.cell(r, 2).value
        created_str = created_val.strftime("%Y-%m-%d %H:%M:%S") if hasattr(created_val, "strftime") else str(created_val)
        topup_rows.append({
            "id": ws_tp.cell(r, 1).value,
            "created": created_str,
            "method": ws_tp.cell(r, 3).value,
            "provider_id": ws_tp.cell(r, 4).value,
            "status": ws_tp.cell(r, 5).value,
            "amount": ws_tp.cell(r, 6).value,
            "currency": ws_tp.cell(r, 7).value,
            "country": ws_tp.cell(r, 8).value,
            "merchant": ws_tp.cell(r, 9).value,
            "merchant_order_id": ws_tp.cell(r, 10).value,
            "note": ws_tp.cell(r, 11).value
        })

with open(os.path.join(FIXTURES_DIR, "topups.json"), "w", encoding="utf-8") as f:
    json.dump(topup_rows, f, indent=2, ensure_ascii=False)
print(f"Saved {len(topup_rows)} topups fixture rows")

wb.close()

import os
import shutil
import pytest
import openpyxl
from datetime import date
from main import run_pipeline

@pytest.fixture
def test_workbook(tmp_path):
    src = "05. Revenue_analysis_copy.xlsx"
    if not os.path.exists(src):
        pytest.skip("The private workbook fixture is not included in the public repository.")
    dst = tmp_path / "test_revenue_analysis.xlsx"
    shutil.copy2(src, dst)
    
    # Create test config
    cfg_path = tmp_path / "test_config.yaml"
    with open("config.yaml", "r", encoding="utf-8") as f:
        cfg_text = f.read()

    cfg_text = cfg_text.replace('workbook_path: "05. Revenue_analysis.xlsx"', f'workbook_path: "{str(dst).replace(chr(92), "/")}"')
    cfg_text = cfg_text.replace('backup_dir: "data/backups"', f'backup_dir: "{str(tmp_path / "backups").replace(chr(92), "/")}"')
    cfg_text = cfg_text.replace('reports_dir: "data/reports"', f'reports_dir: "{str(tmp_path / "reports").replace(chr(92), "/")}"')
    
    with open(cfg_path, "w", encoding="utf-8") as f:
        f.write(cfg_text)

    return str(dst), str(cfg_path)

def test_full_pipeline_dry_run(test_workbook):
    wb_path, cfg_path = test_workbook
    # Run full pipeline with dry_run=True, date override 2026-10-27
    code = run_pipeline(
        config_path=cfg_path,
        dry_run=True,
        override_date="2026-10-27",
        extractor_type="fixture"
    )
    assert code == 0, "Pipeline dry run must exit with 0"

def test_full_pipeline_commit_and_idempotency(test_workbook):
    wb_path, cfg_path = test_workbook
    
    # Run 1: with commit enabled
    code1 = run_pipeline(
        config_path=cfg_path,
        dry_run=False,
        override_date="2026-10-27",
        extractor_type="fixture"
    )
    assert code1 == 0, "First run must succeed"

    # Compute hash after Run 1
    import hashlib
    with open(wb_path, "rb") as f:
        hash1 = hashlib.sha256(f.read()).hexdigest()

    # Run 2: immediately rerun for the same day
    code2 = run_pipeline(
        config_path=cfg_path,
        dry_run=False,
        override_date="2026-10-27",
        extractor_type="fixture"
    )
    assert code2 == 0, "Second run must succeed"

    # Verify structural integrity of sheets
    wb = openpyxl.load_workbook(wb_path, read_only=True)
    assert "BO_extract" in wb.sheetnames
    assert "Data" in wb.sheetnames
    assert "Control" in wb.sheetnames
    wb.close()

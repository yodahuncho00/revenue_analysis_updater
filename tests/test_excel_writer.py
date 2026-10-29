import pytest
from datetime import datetime, date

import openpyxl
from openpyxl.styles import PatternFill

from writers.excel_writer import ExcelWriter, ExcelWriterError


def test_bo_extract_replacement_preserves_existing_cell_formats(tmp_path):
    workbook_path = tmp_path / "bo_extract.xlsx"
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = "BO_extract"
    worksheet["A2"] = datetime(2026, 10, 31)
    worksheet["A2"].number_format = "dd-mmm-yyyy"
    worksheet["H2"] = 1234
    worksheet["H2"].number_format = '_-* # ##0_-;-* # ##0_-;_-* "-"??_-;_-@_-'
    worksheet["H2"].fill = PatternFill(fill_type="solid", fgColor="FFFF00")
    for column in range(1, 28):
        if column not in (1, 8):
            worksheet.cell(2, column).number_format = "General"
    workbook.save(workbook_path)
    workbook.close()

    baseline = openpyxl.load_workbook(workbook_path)
    expected_styles = [baseline["BO_extract"].cell(2, column)._style for column in range(1, 28)]
    baseline.close()

    writer = ExcelWriter(str(workbook_path))
    try:
        row = [None] * 27
        row[0] = datetime(2026, 10, 31)
        row[7] = 1234
        row[23] = '=IF(G{row}="",0,1)'
        writer.update_bo_extract([row, row], date(2026, 10, 31))

        updated = writer.wb["BO_extract"]
        for row_number in (2, 3):
            assert [updated.cell(row_number, column)._style for column in range(1, 28)] == expected_styles
        assert updated["A2"].number_format == "dd-mmm-yyyy"
        assert updated["H2"].number_format == '_-* # ##0_-;-* # ##0_-;_-* "-"??_-;_-@_-'
        assert updated["H2"].fill == PatternFill(fill_type="solid", fgColor="FFFF00")
        assert updated["X2"].value == '=IF(G2="",0,1)'
    finally:
        writer.close()


def test_bo_extract_replacement_skips_unformatted_old_month_rows(tmp_path):
    workbook_path = tmp_path / "bo_extract_unformatted.xlsx"
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = "BO_extract"
    worksheet["A2"] = datetime(2026, 9, 30)
    worksheet["A2"].number_format = "mm-dd-yy"
    worksheet["H2"] = 100
    worksheet["H2"].number_format = '_-* # ##0_-;-* # ##0_-;_-* "-"??_-;_-@_-'
    worksheet["A3"] = datetime(2026, 10, 31)
    worksheet["A3"].number_format = "yyyy-mm-dd h:mm:ss"
    worksheet["H3"] = 200
    expected_date_format = worksheet["A2"].number_format
    expected_amount_format = worksheet["H2"].number_format
    workbook.save(workbook_path)
    workbook.close()

    writer = ExcelWriter(str(workbook_path))
    try:
        row = [None] * 27
        row[0] = datetime(2026, 10, 31)
        row[7] = 1234
        writer.update_bo_extract([row], date(2026, 10, 31))

        updated = writer.wb["BO_extract"]
        assert updated["A3"].number_format == expected_date_format
        assert updated["H3"].number_format == expected_amount_format
    finally:
        writer.close()


def test_bo_extract_replacement_clears_duplicate_rows_missing_period_and_type(tmp_path):
    workbook_path = tmp_path / "bo_extract_orphans.xlsx"
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = "BO_extract"
    rows = []
    for provider in ("", "Paypulse (9772)"):
        row = [None] * 27
        row[0] = datetime(2026, 10, 31)
        row[1] = "M"
        row[2:7] = ["Zimbabwe", "ZWG", "Quicknode", "Ecocash", provider]
        row[7:22] = [15, 0.45, 3, 0.3, 2, 780, 24.96, 3.2, 19.5, 2.5, 0, 0, -790.41, -784.8, 5.61]
        row[23] = '=IF(G{row}="",0,1)'
        rows.append(row)
        for column, value in enumerate(row, 1):
            if column != 24:
                worksheet.cell(len(rows) + 1, column, value)
        worksheet.cell(len(rows) + 1, 1).number_format = "mm-dd-yy"
        worksheet.cell(len(rows) + 1, 8).number_format = "#,##0"
    for offset, row in enumerate(rows):
        orphan_row = 4 + offset
        for column, value in enumerate(row, 1):
            if column not in (1, 2, 3, 24):
                worksheet.cell(orphan_row, column, value)
        worksheet.cell(orphan_row, 3, "Republic of Zimbabwe")
        worksheet.cell(orphan_row, 24, f'=IF(G{orphan_row - 7}="",0,1)')
    workbook.save(workbook_path)
    workbook.close()

    writer = ExcelWriter(str(workbook_path))
    try:
        result = writer.update_bo_extract(rows, date(2026, 10, 31))

        assert result["orphan_rows_cleared"] == 2
        updated = writer.wb["BO_extract"]
        for row_number in (4, 5):
            assert all(updated.cell(row_number, column).value is None for column in range(1, 62))
    finally:
        writer.close()


def test_bo_extract_replacement_rejects_unmatched_rows_missing_period_and_type(tmp_path):
    workbook_path = tmp_path / "bo_extract_unmatched.xlsx"
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = "BO_extract"
    worksheet["A2"] = datetime(2026, 10, 31)
    worksheet["A2"].number_format = "mm-dd-yy"
    worksheet["H2"] = 100
    worksheet["H2"].number_format = "#,##0"
    worksheet["C3"] = "Unmatched existing data"
    worksheet["D3"] = "EUR"
    workbook.save(workbook_path)
    workbook.close()

    writer = ExcelWriter(str(workbook_path))
    try:
        with pytest.raises(ExcelWriterError, match="row 3 has data but no Period or Type"):
            writer.update_bo_extract([[datetime(2026, 10, 31)] + [None] * 26], date(2026, 10, 31))
    finally:
        writer.close()


def test_balance_summary_formula_range_tracks_pasted_data(tmp_path):
    workbook_path = tmp_path / "balance_summary.xlsx"
    workbook = openpyxl.Workbook()
    merchant = workbook.active
    merchant.title = "BO_merchant_balance"
    provider = workbook.create_sheet("BO_provider_balance")
    helper_fill = PatternFill(fill_type="solid", fgColor="FFF8F2E2")
    for worksheet, start, end in ((merchant, 9, 21), (provider, 11, 20)):
        for column in range(start, end + 1):
            worksheet.cell(6, column).fill = helper_fill
    merchant["I5"] = "=1"
    provider["K5"] = "=1"
    merchant["I630"] = "=1"
    provider["K110"] = "=1"
    workbook.save(workbook_path)
    workbook.close()

    writer = ExcelWriter(str(workbook_path))
    try:
        helper_format = '_(* # ##0_);_(* (# ##0);_(* "-"??_);_(@_)'
        merchant_end = writer.update_balance_summary_formulas(
            "BO_merchant_balance",
            "A4",
            622,
            2,
            ["D", "E", "F"],
            "I",
            6,
            ["I", "U"],
            helper_format,
        )
        provider_end = writer.update_balance_summary_formulas(
            "BO_provider_balance",
            "A6",
            96,
            2,
            ["F", "G", "H"],
            "K",
            6,
            ["K", "T"],
            helper_format,
        )

        assert merchant_end == 625
        assert writer.wb["BO_merchant_balance"]["D2"].value == (
            "=SUMIFS(D$6:D$625,$I$6:$I$625,1)"
        )
        assert writer.wb["BO_merchant_balance"]["F2"].value == (
            "=SUMIFS(F$6:F$625,$I$6:$I$625,1)"
        )
        assert provider_end == 101
        assert writer.wb["BO_provider_balance"]["F2"].value == (
            "=SUMIFS(F$6:F$101,$K$6:$K$101,1)"
        )
        assert writer.wb["BO_merchant_balance"]["I5"].number_format == helper_format
        assert writer.wb["BO_merchant_balance"]["U625"].number_format == helper_format
        assert writer.wb["BO_provider_balance"]["K5"].number_format == helper_format
        assert writer.wb["BO_provider_balance"]["T101"].number_format == helper_format
        assert writer.wb["BO_merchant_balance"]["I5"].fill == helper_fill
        assert writer.wb["BO_provider_balance"]["K5"].fill == helper_fill
        assert writer.wb["BO_merchant_balance"]["I5"].value is None
        assert writer.wb["BO_provider_balance"]["K5"].value is None
        assert writer.wb["BO_merchant_balance"]["I630"].value is None
        assert writer.wb["BO_provider_balance"]["K110"].value is None
    finally:
        writer.close()

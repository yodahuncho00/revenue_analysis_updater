import openpyxl
import pytest
from openpyxl.styles import Border, Side

from extractors.base import BalanceTableCapture
from writers.excel_table_paster import ExcelTablePaster


def test_plain_text_fallback_keeps_multiline_cell_content_in_one_row():
    capture = BalanceTableCapture(
        rows=[
            {
                "row_class": "spacer-row",
                "cells": [""],
                "spans": [{"col_span": 8, "row_span": 1}],
            },
            {
                "row_class": "account-row",
                "cells": [
                    "",
                    "Zoykpay Tembo",
                    "748 Tanzania\n\n749 Tanzania\n\n750 Tanzania",
                    "",
                    "Oct 29, 12:00 PM",
                    "30,382,668.77",
                    "30,382,668.77",
                    "16,444,236.88",
                ],
            },
        ],
        html="<table></table>",
        row_count=2,
        column_count=8,
    )

    plain_text = ExcelTablePaster._plain_text(capture)
    lines = plain_text.splitlines()

    assert len(lines) == capture.row_count
    assert lines[0] == "\t".join([""] * 8)
    assert lines[1].split("\t")[2] == "748 Tanzania 749 Tanzania 750 Tanzania"
    assert lines[1].split("\t")[5:] == [
        "30,382,668.77",
        "30,382,668.77",
        "16,444,236.88",
    ]


def test_plain_text_fallback_places_rowspan_continuation_in_next_free_column():
    capture = BalanceTableCapture(
        rows=[
            {
                "row_class": "account-row",
                "cells": ["Provider", "First account"],
                "spans": [
                    {"col_span": 1, "row_span": 2},
                    {"col_span": 1, "row_span": 1},
                ],
            },
            {
                "row_class": "account-row",
                "cells": ["Second account"],
                "spans": [{"col_span": 1, "row_span": 1}],
            },
        ],
        html="<table></table>",
        row_count=2,
        column_count=2,
    )

    lines = ExcelTablePaster._plain_text(capture).splitlines()

    assert lines[0].split("\t") == ["Provider", "First account"]
    assert lines[1].split("\t") == ["", "Second account"]


def test_provider_paste_keeps_multiline_accounts_in_cell_and_separator_aligned(
    tmp_path,
):
    pytest.importorskip("win32com.client")
    workbook_path = tmp_path / "provider_multiline.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active.title = "BO_merchant_balance"
    workbook.create_sheet("BO_provider_balance")
    workbook.save(workbook_path)
    workbook.close()

    capture = BalanceTableCapture(
        rows=[
            {
                "row_class": "account-row",
                "height": 60,
                "cells": [
                    "",
                    "Zoykpay Tembo",
                    "748 Tanzania\n\n749 Tanzania",
                    "",
                    "Oct 29, 1:00 PM",
                    "30,382,668.77",
                    "30,382,668.77",
                    "16,444,236.88",
                ],
            },
            {
                "row_class": "spacer-row",
                "height": 28,
                "cells": [""],
                "spans": [{"col_span": 8, "row_span": 1}],
            },
            {
                "row_class": "country-row",
                "height": 28,
                "cells": ["TZS", "", "", "", "", "30,382,668.77",
                          "30,382,668.77", "16,444,236.88"],
            },
        ],
        html=(
            "<table><tbody>"
            "<tr style='height:60px'><td></td><td>Zoykpay Tembo</td>"
            "<td>748 Tanzania<br><br>749 Tanzania</td><td></td>"
            "<td>Oct 29, 1:00 PM</td><td>30,382,668.77</td>"
            "<td>30,382,668.77</td><td>16,444,236.88</td></tr>"
            "<tr style='height:28px'><td colspan='8'></td></tr>"
            "<tr style='height:28px'><td>TZS</td><td></td><td></td><td></td>"
            "<td></td><td>30,382,668.77</td><td>30,382,668.77</td>"
            "<td>16,444,236.88</td></tr>"
            "</tbody></table>"
        ),
        row_count=3,
        column_count=8,
    )

    ExcelTablePaster.paste_tables(
        str(workbook_path),
        {"BO_provider_balance": (capture, "A6")},
    )

    updated = openpyxl.load_workbook(workbook_path, data_only=True)
    pasted = updated["BO_provider_balance"]
    assert pasted["C6"].value == "748 Tanzania\n\n749 Tanzania"
    assert pasted["C6"].alignment.wrap_text is True
    assert all(pasted.cell(7, column).value is None for column in range(1, 9))
    assert pasted["A8"].value == "TZS"
    assert pasted["F8"].value == pytest.approx(30382668.77)
    assert pasted.row_dimensions[6].height == pytest.approx(45)
    assert pasted.row_dimensions[7].height == pytest.approx(21)
    updated.close()


def test_pastes_complete_html_table_and_unmerges_without_touching_adjacent_cells(tmp_path):
    pytest.importorskip("win32com.client")
    workbook_path = tmp_path / "balance_table.xlsx"
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = "BO_merchant_balance"
    worksheet["G4"] = "untouched"
    worksheet["I4"] = "=SUM(1,2)"
    worksheet["I4"].number_format = "#,##0.00"
    workbook.create_sheet("BO_provider_balance")
    workbook.save(workbook_path)
    workbook.close()

    capture = BalanceTableCapture(
        rows=[
            {"row_class": "country-row", "cells": ["Country", "Currency"]},
            {"row_class": "spacer-row", "cells": ["", ""]},
            {
                "row_class": "account-row",
                "cells": ["Merchant", "First method"],
                "spans": [
                    {"col_span": 1, "row_span": 2},
                    {"col_span": 1, "row_span": 1},
                ],
            },
            {
                "row_class": "account-row",
                "cells": ["Second method"],
                "spans": [{"col_span": 1, "row_span": 1}],
            },
        ],
        html=(
            '<table><tbody>'
            '<tr><td style="background-color: rgb(255, 255, 0); font-weight: bold">'
            'Country</td><td>Currency</td></tr>'
            '<tr><td></td><td></td></tr>'
            '<tr><td rowspan="2">Merchant</td><td>First method</td></tr>'
            '<tr><td>Second method</td></tr>'
            '</tbody></table>'
        ),
        row_count=4,
        column_count=2,
    )

    ExcelTablePaster.paste_tables(
        str(workbook_path),
        {"BO_merchant_balance": (capture, "A4")},
    )

    updated = openpyxl.load_workbook(workbook_path, data_only=False)
    pasted = updated["BO_merchant_balance"]
    assert pasted["A4"].value == "Country"
    assert pasted["B4"].value == "Currency"
    assert pasted["A5"].value is None
    assert pasted["B5"].value is None
    assert pasted["A6"].value == "Merchant"
    assert pasted["B6"].value == "First method"
    assert pasted["A7"].value is None
    assert pasted["B7"].value == "Second method"
    assert pasted["G4"].value == "untouched"
    assert pasted["I4"].value == "=SUM(1,2)"
    assert pasted["I4"].number_format == "#,##0.00"
    assert pasted["A4"].font.bold is True
    assert pasted["A4"].fill.fgColor.rgb in ("FFFFFF00", "00FFFF00")
    assert all(pasted.row_dimensions[row].height == 13 for row in range(4, 8))
    assert not list(pasted.merged_cells.ranges)
    updated.close()


def test_pastes_amounts_as_numbers_with_expected_locale_format(tmp_path):
    pytest.importorskip("win32com.client")
    workbook_path = tmp_path / "balance_amounts.xlsx"
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = "BO_merchant_balance"
    worksheet["F7"] = "stale tail"
    worksheet["F7"].border = Border(bottom=Side(style="thin"))
    workbook.create_sheet("BO_provider_balance")
    workbook.save(workbook_path)
    workbook.close()

    capture = BalanceTableCapture(
        rows=[
            {"row_class": "country-row", "cells": ["Zambia", "ZMW", "", "165,090,263.26", "0.00", "165,090,263.26"]},
            {"row_class": "account-row", "cells": ["", "Bank", "Account", "261 745 919,51", "—", "261 745 919,51"]},
        ],
        html=(
            "<table><tbody>"
            "<tr><td>Zambia</td><td>ZMW</td><td></td>"
            "<td>165,090,263.26</td><td>0.00</td><td>165,090,263.26</td></tr>"
            "<tr><td></td><td>Bank</td><td>Account</td>"
            "<td>261 745 919,51</td><td>—</td><td>261 745 919,51</td></tr>"
            "</tbody></table>"
        ),
        row_count=2,
        column_count=6,
    )

    ExcelTablePaster.paste_tables(
        str(workbook_path),
        {"BO_merchant_balance": (capture, "A4")},
    )

    updated = openpyxl.load_workbook(workbook_path, data_only=False)
    pasted = updated["BO_merchant_balance"]
    assert pasted["D4"].value == pytest.approx(165090263.26)
    assert pasted["D5"].value == pytest.approx(261745919.51)
    assert pasted["E5"].value == "—"
    assert pasted["F5"].value == pytest.approx(261745919.51)
    assert pasted["F7"].value is None
    assert pasted["F7"].border.bottom.style is None
    assert pasted["D4"].number_format == "#,##0.00;\\-#,##0.00;0"
    assert all(pasted.row_dimensions[row].height == 13 for row in (4, 5, 6, 7))
    updated.close()

    import os
    import win32com.client

    excel = win32com.client.DispatchEx("Excel.Application")
    excel.Visible = False
    excel.DisplayAlerts = False
    try:
        displayed_workbook = excel.Workbooks.Open(
            os.path.abspath(workbook_path), UpdateLinks=0, ReadOnly=True
        )
        assert (
            displayed_workbook.Worksheets("BO_merchant_balance").Range("D5").Text
            == "261 745 919,51"
        )
        assert (
            displayed_workbook.Worksheets("BO_merchant_balance").Range("E4").Text
            == "0"
        )
        displayed_workbook.Close(SaveChanges=False)
    finally:
        excel.Quit()


def test_provider_paste_skips_header_when_normalizing_amounts(tmp_path):
    pytest.importorskip("win32com.client")
    workbook_path = tmp_path / "provider_amounts.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active.title = "BO_merchant_balance"
    provider = workbook.create_sheet("BO_provider_balance")
    provider["A5"] = "Provider"
    provider["H5"] = "Merchants Balance"
    provider["K5"] = "=1"
    provider["T5"] = '=\"old header\"'
    provider["N7"] = 44933182.8
    provider["C10"] = "stale provider row"
    provider["C10"].border = Border(bottom=Side(style="thin"))
    provider.merge_cells("A11:H12")
    provider["A11"] = "stale merged provider row"
    provider.merge_cells("A13:H14")
    workbook.save(workbook_path)
    workbook.close()

    capture = BalanceTableCapture(
        rows=[
            {
                "row_class": "country-row",
                "height": 30,
                "cells": ["CDF", "", "", "", "", "0.00", "0.00", "0.00"],
            },
            {
                "row_class": "account-row",
                "height": 60,
                "cells": ["", "Account", "ID", "refresh", "Oct 29, 8:00 AM",
                          "261,745,919.51", "204,322,189.57", "165,604,521.80"],
            },
        ],
        html=(
            "<table><tbody><tr class='country-row' style='height: 30px'>"
            "<td style='text-align: left'>CDF</td><td></td>"
            "<td></td><td></td><td></td><td>0.00</td><td>0.00</td><td>0.00</td></tr>"
            "<tr class='account-row' style='height: 60px'><td></td><td>Account</td>"
            "<td>ID</td><td>refresh</td><td>Oct 29, 8:00 AM</td>"
            "<td>261,745,919.51</td><td>204,322,189.57</td>"
            "<td>165,604,521.80</td></tr></tbody></table>"
        ),
        row_count=2,
        column_count=8,
    )

    ExcelTablePaster.paste_tables(
        str(workbook_path),
        {"BO_provider_balance": (capture, "A6")},
    )

    updated = openpyxl.load_workbook(workbook_path, data_only=False)
    pasted = updated["BO_provider_balance"]
    assert all(pasted.cell(5, column).value is None for column in range(1, 9))
    assert all(pasted.cell(5, column).value is None for column in range(11, 21))
    assert pasted["A6"].value == "CDF"
    assert pasted["F6"].value == 0
    assert pasted["F7"].value == pytest.approx(261745919.51)
    assert pasted["G7"].value == pytest.approx(204322189.57)
    assert pasted["H7"].value == pytest.approx(165604521.80)
    assert pasted["C10"].value is None
    assert pasted["C10"].border.bottom.style is None
    assert pasted["A11"].value is None
    assert not list(pasted.merged_cells.ranges)
    assert pasted["E7"].value == "Oct 29, 8:00 AM"
    assert pasted["F7"].number_format == "#,##0.00;\\-#,##0.00;0"
    assert pasted["A6"].alignment.horizontal == "left"
    assert pasted["A7"].alignment.horizontal in (None, "general")
    assert pasted.row_dimensions[6].height == pytest.approx(22.5)
    assert pasted.row_dimensions[7].height == pytest.approx(45)
    updated.close()

    import os
    import win32com.client

    excel = win32com.client.DispatchEx("Excel.Application")
    excel.Visible = False
    excel.DisplayAlerts = False
    try:
        displayed_workbook = excel.Workbooks.Open(
            os.path.abspath(workbook_path), UpdateLinks=0, ReadOnly=True
        )
        displayed_provider = displayed_workbook.Worksheets("BO_provider_balance")
        assert displayed_provider.Range("F6").Text == "0"
        assert displayed_provider.Range("N7").Text.strip() == "44 933 183"
        displayed_workbook.Close(SaveChanges=False)
    finally:
        excel.Quit()


def test_amount_parser_handles_grouping_and_decimal_comma():
    assert ExcelTablePaster._parse_amount("165,090,263.26") == pytest.approx(165090263.26)
    assert ExcelTablePaster._parse_amount("261 745 919,51") == pytest.approx(261745919.51)
    assert ExcelTablePaster._parse_amount("—") == "—"
    with pytest.raises(ValueError, match="Unexpected balance amount"):
        ExcelTablePaster._parse_amount("Oct 29, 8:00 AM")


def test_retries_excel_call_rejected_errors(monkeypatch):
    import pywintypes

    monkeypatch.setattr("writers.excel_table_paster.time.sleep", lambda _: None)
    calls = 0

    def operation():
        nonlocal calls
        calls += 1
        if calls < 3:
            raise pywintypes.com_error(
                ExcelTablePaster.RPC_E_CALL_REJECTED,
                "Call was rejected by the callee.",
                None,
                None,
            )
        return "complete"

    assert ExcelTablePaster._retry_excel_call(operation) == "complete"
    assert calls == 3

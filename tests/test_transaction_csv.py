import csv
import io
from datetime import datetime

import openpyxl
import pytest

from extractors.playwright_extractor import PlaywrightExtractionError, PlaywrightExtractor
from transformers.append_only import AppendOnlyTransformer
from writers.excel_writer import ExcelWriter, TRANSACTION_DATE_NUMBER_FORMAT


def _csv_content(rows):
    output = io.StringIO(newline="")
    writer = csv.DictWriter(output, fieldnames=PlaywrightExtractor.TRANSACTION_CSV_COLUMNS)
    writer.writeheader()
    writer.writerows(rows)
    return output.getvalue().encode("utf-8-sig")


def test_transaction_csv_export_supplies_status_country_and_full_timestamp():
    content = _csv_content([{
        "Id": "12345",
        "Created": "2026-10-28 11:37:18.139839",
        "Method": "Mobile Money",
        "Provider Id": "289",
        "Status": "Success",
        "Amount": "1250.50",
        "Currency": "ZMW",
        "Country": "ZMB",
        "Merchant": "Example Merchant",
        "Merchant Order Id": "order-1",
        "Note": "settlement reference",
    }])

    records = PlaywrightExtractor._parse_transaction_csv(
        content,
        [{"id": 12345}],
        "settlements",
    )
    row = AppendOnlyTransformer.transform_records(records, set(), "Settlements")[0]

    assert row[1] == datetime(2026, 10, 28, 11, 37, 18, 139839)
    assert row[4] == "Success"
    assert row[7] == "ZMB"
    assert row[9] == "order-1"
    assert row[10] == "settlement reference"


def test_transaction_csv_export_must_match_api_ids():
    content = _csv_content([{
        "Id": "12345",
        "Created": "2026-10-28 11:37:18",
        "Method": "",
        "Provider Id": "",
        "Status": "Success",
        "Amount": "1",
        "Currency": "ZMW",
        "Country": "ZMB",
        "Merchant": "",
        "Merchant Order Id": "",
        "Note": "",
    }])

    with pytest.raises(PlaywrightExtractionError, match="absent from the filtered API"):
        PlaywrightExtractor._parse_transaction_csv(
            content,
            [{"id": 99999}],
            "topups",
        )


def test_transaction_csv_export_allows_api_state_3_records_omitted_from_csv():
    content = _csv_content([{
        "Id": "12345",
        "Created": "2026-10-28 11:37:18",
        "Method": "Mobile Money",
        "Provider Id": "289",
        "Status": "Success",
        "Amount": "1",
        "Currency": "ZMW",
        "Country": "ZMB",
        "Merchant": "Example Merchant",
        "Merchant Order Id": "order-1",
        "Note": "reference",
    }])

    records = PlaywrightExtractor._parse_transaction_csv(
        content,
        [{"id": 12345}, {"id": 99999, "state": 3}],
        "settlements",
    )

    assert [record["id"] for record in records] == ["12345"]
    assert records[0]["merchant_order_id"] == "order-1"
    assert records[0]["note"] == "reference"


def test_transaction_csv_export_rejects_non_state_3_api_omissions():
    content = _csv_content([{
        "Id": "12345",
        "Created": "2026-10-28 11:37:18",
        "Method": "Mobile Money",
        "Provider Id": "289",
        "Status": "Success",
        "Amount": "1",
        "Currency": "ZMW",
        "Country": "ZMB",
        "Merchant": "Example Merchant",
        "Merchant Order Id": "order-1",
        "Note": "reference",
    }])

    with pytest.raises(PlaywrightExtractionError, match="not in excluded state 3"):
        PlaywrightExtractor._parse_transaction_csv(
            content,
            [{"id": 12345}, {"id": 99999, "state": 2}],
            "settlements",
        )


@pytest.mark.parametrize("sheet_name,append_method", [
    ("BO_settlements", "append_settlements"),
    ("BO_topups", "append_topups"),
])
def test_transaction_writer_applies_localized_datetime_format(
    tmp_path, sheet_name, append_method
):
    assert TRANSACTION_DATE_NUMBER_FORMAT == "dd.mm.yyyy"
    workbook_path = tmp_path / "transactions.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active.title = sheet_name
    workbook.create_sheet(
        "BO_topups" if sheet_name == "BO_settlements" else "BO_settlements"
    )
    worksheet = workbook[sheet_name]
    worksheet["A2"] = 1
    worksheet["B2"] = datetime(2026, 10, 27, 9, 0)
    worksheet["B2"].number_format = "mm-dd-yy"
    workbook.save(workbook_path)
    workbook.close()

    writer = ExcelWriter(str(workbook_path))
    try:
        new_row = [2, datetime(2026, 10, 28, 11, 37, 18)] + [None] * 11
        getattr(writer, append_method)([new_row])

        updated = writer.wb[sheet_name]
        assert updated["B2"].number_format == TRANSACTION_DATE_NUMBER_FORMAT
        assert updated["B3"].number_format == TRANSACTION_DATE_NUMBER_FORMAT
        assert updated["B3"].value == datetime(2026, 10, 28, 11, 37, 18)
        assert updated["B3"].value.time().hour == 11
    finally:
        writer.close()


@pytest.mark.parametrize("sheet_name", ["BO_settlements", "BO_topups"])
def test_transaction_detail_backfill_only_fills_empty_j_and_k(tmp_path, sheet_name):
    workbook_path = tmp_path / "transactions_backfill.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active.title = sheet_name
    workbook.create_sheet(
        "BO_topups" if sheet_name == "BO_settlements" else "BO_settlements"
    )
    worksheet = workbook[sheet_name]
    worksheet.append(["Id", "Created", "Method", "Provider Id", "Status",
                      "Amount", "Currency", "Country", "Merchant",
                      "Merchant Order Id", "Note"])
    worksheet.append([12345, datetime(2026, 10, 28, 11, 37), None, None,
                      None, None, None, None, None, None, None])
    worksheet.append([12346, datetime(2026, 10, 28, 12, 0), None, None,
                      None, None, None, None, None, "existing-order", ""])
    workbook.save(workbook_path)
    workbook.close()

    writer = ExcelWriter(str(workbook_path))
    try:
        updated = writer.backfill_transaction_details(sheet_name, [
            {"id": "12345", "merchant_order_id": "downloaded-order",
             "note": "downloaded-note"},
            {"id": "12346", "merchant_order_id": "replacement-order",
             "note": "downloaded-note-2"},
        ])

        assert updated == {"merchant_order_id": 1, "note": 2}
        assert writer.wb[sheet_name]["J2"].value == "downloaded-order"
        assert writer.wb[sheet_name]["K2"].value == "downloaded-note"
        assert writer.wb[sheet_name]["J3"].value == "existing-order"
        assert writer.wb[sheet_name]["K3"].value == "downloaded-note-2"
    finally:
        writer.close()
